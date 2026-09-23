from io import BytesIO
from pathlib import Path
from xml.sax.saxutils import escape
import os
from .schemas import Meeting


def stamp(seconds: float) -> str:
    seconds = int(seconds)
    return f'{seconds // 3600:02d}:{seconds // 60 % 60:02d}:{seconds % 60:02d}'


def task_lines(task):
    due = task.due_date.isoformat() if task.due_date else task.deadline_text or 'Не указан'
    status = {'open': 'Открыто', 'in_progress': 'В работе', 'done': 'Выполнено'}[task.status]
    lines = [f'Ответственный: {task.assignee_name or "Не установлен"}', f'Срок: {due}', f'Статус: {status}']
    if task.is_overdue:
        lines.append('Просрочено')
    if task.due_date and task.deadline_text:
        lines.append(f'Исходная формулировка срока: {task.deadline_text}')
    if task.needs_review:
        lines.append('Требует уточнения: ' + '; '.join(task.review_reasons))
    return lines


def meeting_date(meeting):
    from zoneinfo import ZoneInfo
    return (meeting.started_at.astimezone(ZoneInfo(meeting.timezone)).strftime('%d.%m.%Y %H:%M')
            if meeting.started_at else 'Дата совещания не указана')


def font_paths():
    candidates = [
        Path(os.environ.get('HACKALEM_FONT_DIR', '/usr/share/fonts/truetype/dejavu')),
        Path('C:/Windows/Fonts'),
        Path('/usr/share/fonts/truetype/liberation2'),
    ]
    for directory in candidates:
        for normal, bold in [('DejaVuSans.ttf', 'DejaVuSans-Bold.ttf'),
                             ('arial.ttf', 'arialbd.ttf'), ('LiberationSans-Regular.ttf', 'LiberationSans-Bold.ttf')]:
            if (directory / normal).is_file() and (directory / bold).is_file():
                return str(directory / normal), str(directory / bold)
    raise RuntimeError('Install fonts-dejavu-core or set HACKALEM_FONT_DIR to a directory with DejaVuSans fonts')


def export_pdf(meeting: Meeting) -> bytes:
    from reportlab.lib import colors
    from reportlab.lib.enums import TA_LEFT
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import mm
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, PageBreak, LongTable, TableStyle, KeepTogether
    regular, bold = font_paths()
    if 'MeetingSans' not in pdfmetrics.getRegisteredFontNames():
        pdfmetrics.registerFont(TTFont('MeetingSans', regular))
        pdfmetrics.registerFont(TTFont('MeetingSansBold', bold))
    base = ParagraphStyle('Body', fontName='MeetingSans', fontSize=10, leading=15,
                          spaceAfter=6, alignment=TA_LEFT, splitLongWords=True)
    heading = ParagraphStyle('Heading', parent=base, fontName='MeetingSansBold', fontSize=13,
                             leading=18, spaceBefore=14, spaceAfter=8, keepWithNext=True)
    title = ParagraphStyle('Title', parent=heading, fontSize=20, leading=26, spaceBefore=0)
    muted = ParagraphStyle('Muted', parent=base, fontSize=9, leading=13, textColor=colors.HexColor('#475569'))
    cell = ParagraphStyle('Cell', parent=base, fontSize=8.5, leading=11.5, spaceAfter=3)
    cell_muted = ParagraphStyle('CellMuted', parent=cell, fontSize=8, leading=10.5, textColor=colors.HexColor('#475569'))
    transcript_heading = ParagraphStyle('TranscriptHeading', parent=heading, fontSize=10, leading=13, spaceBefore=8, spaceAfter=3)
    story = []
    def add(text, style=base):
        story.append(Paragraph(escape(str(text)).replace('\n', '<br/>'), style))
    add('Протокол совещания', title)
    add(meeting.title, heading)
    add(f'{meeting_date(meeting)} · {meeting.timezone}', muted)
    add('Проект подготовлен ИИ. Поручения и неоднозначные сведения подлежат проверке секретарём.', muted)
    add('Краткое содержание', heading)
    add(meeting.summary.overview or 'Содержание не сформировано.')
    for caption, items in [('Риски и открытые вопросы', meeting.summary.risks)]:
        if items:
            add(caption, heading)
            for item in items:
                add('• ' + item)
    segment_map = {segment.id: segment for segment in meeting.segments}
    if not meeting.tasks:
        add('Поручения (0)', heading)
        add('Поручения в записи не обнаружены.')
    if meeting.tasks:
        rows = [[Paragraph(f'Поручения ({len(meeting.tasks)})', heading), '', '', ''],
                [Paragraph(f'<b>{label}</b>', cell) for label in ('№', 'Поручение и источник', 'Ответственный', 'Срок / статус')]]
        for index, task in enumerate(meeting.tasks, 1):
            action = [Paragraph('<b>' + escape(task.title) + '</b>', cell)]
            if task.description and task.description.strip().rstrip('.') != task.title.strip().rstrip('.'):
                action.append(Paragraph(escape(task.description), cell))
            evidence = [segment_map[sid] for sid in task.evidence_segment_ids if sid in segment_map]
            if evidence:
                action.append(Paragraph('Источники: ' + ', '.join(f'{stamp(s.start)}–{stamp(s.end)}' for s in evidence), cell_muted))
            if task.needs_review:
                action.append(Paragraph('Уточнить: ' + escape('; '.join(task.review_reasons)), cell_muted))
            deadline = task.due_date.isoformat() if task.due_date else task.deadline_text or 'Не указан'
            state = {'open': 'Открыто', 'in_progress': 'В работе', 'done': 'Выполнено'}[task.status]
            dates = [Paragraph(escape(deadline), cell), Paragraph(state, cell_muted)]
            if task.due_date and task.deadline_text:
                dates.append(Paragraph('В записи: ' + escape(task.deadline_text), cell_muted))
            if task.is_overdue:
                dates.append(Paragraph('<b>Просрочено</b>', cell))
            rows.append([Paragraph(str(index), cell), action,
                         Paragraph(escape(task.assignee_name or 'Не установлен'), cell), dates])
        table = LongTable(rows, colWidths=[8 * mm, 91 * mm, 35 * mm, 36 * mm], repeatRows=2,
                          splitInRow=1, hAlign='LEFT')
        table.setStyle(TableStyle([
            ('SPAN', (0, 0), (-1, 0)),
            ('BACKGROUND', (0, 1), (-1, 1), colors.HexColor('#e2e8f0')),
            ('GRID', (0, 1), (-1, -1), .4, colors.HexColor('#cbd5e1')),
            ('VALIGN', (0, 0), (-1, -1), 'TOP'),
            ('LEFTPADDING', (0, 0), (-1, -1), 5), ('RIGHTPADDING', (0, 0), (-1, -1), 5),
            ('TOPPADDING', (0, 0), (-1, -1), 6), ('BOTTOMPADDING', (0, 0), (-1, -1), 6),
        ]))
        story.append(table)
    story.append(PageBreak())
    add('Транскрипт', heading)
    names = {speaker.id: speaker.display_name for speaker in meeting.speakers}
    for segment in meeting.segments:
        block = [Paragraph(escape(f'{stamp(segment.start)}–{stamp(segment.end)} · {names.get(segment.speaker_id, "Неизвестный спикер")}'), transcript_heading),
                 Paragraph(escape(segment.text), base)]
        if segment.needs_review:
            block.append(Paragraph('Распознавание или определение говорящего требует проверки.', muted))
        story.append(KeepTogether(block))
    stream = BytesIO()
    document = SimpleDocTemplate(stream, pagesize=A4, leftMargin=18 * mm, rightMargin=18 * mm,
                                 topMargin=18 * mm, bottomMargin=20 * mm,
                                 title=meeting.title, author='HackAlem AI')
    def footer(canvas, doc):
        canvas.saveState()
        canvas.setFont('MeetingSans', 8)
        canvas.setFillColor(colors.HexColor('#64748b'))
        canvas.drawString(18 * mm, 10 * mm, 'Проект протокола')
        canvas.drawRightString(A4[0] - 18 * mm, 10 * mm, str(doc.page))
        canvas.restoreState()
    document.build(story, onFirstPage=footer, onLaterPages=footer)
    return stream.getvalue()


def export_docx(meeting: Meeting) -> bytes:
    from docx import Document
    from docx.shared import Pt, Mm, RGBColor
    document = Document()
    section = document.sections[0]
    section.page_width, section.page_height = Mm(210), Mm(297)
    section.top_margin = section.bottom_margin = Mm(18)
    section.left_margin = section.right_margin = Mm(18)
    normal = document.styles['Normal']
    normal.font.name, normal.font.size = 'Arial', Pt(10)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.15
    for name in ('Title', 'Heading 1', 'Heading 2'):
        document.styles[name].font.name = 'Arial'
        document.styles[name].font.color.rgb = RGBColor(0, 0, 0)
    document.add_paragraph('Протокол совещания', 'Title')
    document.add_paragraph(meeting.title, 'Heading 1')
    document.add_paragraph(f'{meeting_date(meeting)} · {meeting.timezone}')
    document.add_paragraph('Проект подготовлен ИИ. Поручения и неоднозначные сведения подлежат проверке секретарём.')
    document.add_heading('Краткое содержание', level=1)
    document.add_paragraph(meeting.summary.overview or 'Содержание не сформировано.')
    for caption, items in [('Решения', meeting.summary.decisions), ('Риски и открытые вопросы', meeting.summary.risks)]:
        if items:
            document.add_heading(caption, level=1)
            for item in items:
                document.add_paragraph(item, style='List Bullet')
    document.add_heading('Поручения', level=1)
    if not meeting.tasks:
        document.add_paragraph('Поручения в записи не обнаружены.')
    segment_map = {segment.id: segment for segment in meeting.segments}
    for index, task in enumerate(meeting.tasks, 1):
        document.add_heading(f'{index}. {task.title}', level=2)
        if task.description:
            document.add_paragraph(task.description)
        for line in task_lines(task):
            document.add_paragraph(line)
        sources = [segment_map[sid] for sid in task.evidence_segment_ids if sid in segment_map]
        if sources:
            document.add_paragraph('Источники: ' + ', '.join(f'{stamp(s.start)}–{stamp(s.end)}' for s in sources))
    document.add_page_break()
    document.add_heading('Транскрипт', level=1)
    names = {speaker.id: speaker.display_name for speaker in meeting.speakers}
    for segment in meeting.segments:
        document.add_heading(f'{stamp(segment.start)}–{stamp(segment.end)} · {names.get(segment.speaker_id, "Неизвестный спикер")}', level=2)
        document.add_paragraph(segment.text)
        if segment.needs_review:
            document.add_paragraph('Распознавание или определение говорящего требует проверки.')
    document.core_properties.author = 'HackAlem AI'
    document.core_properties.title = meeting.title
    stream = BytesIO()
    document.save(stream)
    return stream.getvalue()
